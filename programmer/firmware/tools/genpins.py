import json, re, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
SP='/tmp/claude-1000/-home-frank-data-projects-bloxorz/213a7a2c-60a8-4f2e-8fc4-941c36452dd9/scratchpad/'
d=json.load(open(SP+'pindata.json'))
u3, cart, con, nets = d['u3'], d['cart'], d['con'], d['nets']
DIRS={'A':'out','D':'bidir','CE':'out','OE':'out','RW':'out','HALT':'out','NMI':'out','CART':'out',
      'PB6':'bidir','DIR_D':'out','DIR_PB6':'out','SHIFT_OE':'out'}
def cart_dir(sig):
    if re.fullmatch(r'A\d+',sig): return 'out'
    if re.fullmatch(r'D\d+',sig): return 'bidir'
    return DIRS.get(sig,'')
OTHER={'USB_DM':('USB D-','bidir'),'USB_DP':('USB D+','bidir'),'SWCLK':('SWD clock','in'),
       'SWDIO':('SWD data','bidir'),'RUN':('reset (Tag-Connect)','in'),'XIN':('12 MHz oscillator','in'),
       'XOUT':('test point J2','')}
CODEC_DIR={'MCLK':'out','BCLK':'out','WCLK':'out','DIN':'out','DOUT':'in','SCL':'out','SDA':'bidir',
           '~{RESET}':'out'}
wb=openpyxl.Workbook(); ws=wb.active; ws.title='RP2354B pins'
hdr=['RP2350 pin','Pin name','GPIO','Net','Dir','Vectrex pin','Vectrex signal','Level shifter',
     'Codec (U7)','Flash','Other','Notes']
ws.append(hdr)
def gpio_note(g, net):
    n=[]
    if g is None: return ''
    if g==0: n.append('QMI CS1 capable (GPIO0/8/19/47) - external flash CS')
    if 1<=g<=31: n.append('cartridge PIO window, GPIOBASE 0 (GPIO0-31)')
    if g in (40,41): n.append('I2C0 pair (SDA=40, SCL=41)')
    if 42<=g<=46: n.append('I2S via PIO, GPIOBASE 16; BCLK/WCLK adjacent for side-set')
    return '; '.join(n)
rows=[]
for pin,(name,net) in u3.items():
    try: pn=int(pin)
    except ValueError: continue
    g=None
    m=re.match(r'GPIO(\d+)',name)
    if m: g=int(m.group(1))
    vec_pin=vec_sig=shift=codec=flash=other=direction=''
    if net.startswith('/RP_') and net in cart:
        cnet,ref,ch=cart[net]; sig=cnet.split('/')[-1]
        vec_sig=sig; shift='%s %s'%(ref,ch)
        if cnet in con: vec_pin=con[cnet][0]
        direction=cart_dir(sig)
    partners=[p for p in nets.get(net,[]) if p[0]!='U3']
    for ref,pp in partners:
        if ref=='U7':
            for pin2,(nm,_) in []: pass
    # codec / flash / other from the net name
    s=net.lstrip('/')
    if s.startswith('I2S_') or s in ('SDA','SCL','CODEC_RESET'):
        f={'I2S_MCLK':'MCLK','I2S_BCLK':'BCLK','I2S_WCLK':'WCLK','I2S_DIN':'DIN','I2S_DOUT':'DOUT',
           'SDA':'SDA','SCL':'SCL','CODEC_RESET':'~RESET'}.get(s,s)
        codec=f; direction=CODEC_DIR.get(f if f!='~RESET' else '~{RESET}','')
    if s.startswith('QSPI_') or s.startswith('FLASH_') or name.startswith('QSPI'):
        flash={'FLASH_CS':'CS (bank 0-3 via 138)','FLASH_SEL0':'bank select 0','FLASH_SEL1':'bank select 1'}.get(s, name)
        if not direction: direction='bidir' if 'SD' in name else 'out'
    if s=='SPK_SHUTDOWN': other='speaker amp shutdown (U14/U15)'; direction='out'
    if name in OTHER: other,direction=OTHER[name]
    if net.startswith('Net-(U3-GPIO39)'): other='status LED'; direction='out'
    if net.startswith('Net-(J3'): other='test point J3'
    if net.startswith('Net-(J4'): other='test point J4'
    if name in ('IOVDD','DVDD','VREG_VIN','VREG_AVDD','ADC_AVDD','USB_OTP_VDD','QSPI_IOVDD','VREG_PGND','VREG_FB','VREG_LX','GND'):
        other='power'; direction=''
    rows.append([pn,name,g if g is not None else '',net if not net.startswith('unconnected') else '(not connected)',
                 direction,vec_pin,vec_sig,shift,codec,flash,other,gpio_note(g,net)])
for r in sorted(rows): ws.append(r)
ws.freeze_panes='A2'
for c in range(1,len(hdr)+1):
    ws.cell(1,c).font=Font(bold=True)
    ws.cell(1,c).fill=PatternFill('solid',fgColor='DDDDDD')
    w=max(len(str(ws.cell(r,c).value or '')) for r in range(1,ws.max_row+1))
    ws.column_dimensions[get_column_letter(c)].width=min(max(w+2,9),46)
ws.auto_filter.ref=ws.dimensions
wb.save('/home/frank/data/projects/bloxorz/programmer/pins.xlsx')
print('pins.xlsx written: %d pins'%(ws.max_row-1))
